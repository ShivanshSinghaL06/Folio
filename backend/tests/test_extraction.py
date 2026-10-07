import json
from io import BytesIO

import pytest
from docx import Document

from app.core.errors import AppError
from app.services.embeddings.bedrock import titan_embed_request
from app.services.generation.bedrock import _response_text
from app.services.ingestion.extract import extract_document
from app.services.uploads import validate_upload


def test_docx_keeps_headings_and_table_text():
    document = Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Employees receive twenty days.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Carryover is five days."
    buffer = BytesIO()
    document.save(buffer)

    blocks = extract_document("leave policy.docx", buffer.getvalue())
    assert any(block.section == "Leave policy" and "twenty days" in block.text for block in blocks)
    assert any("Carryover" in block.text for block in blocks)


def test_unreadable_pdf_is_rejected():
    with pytest.raises(AppError) as caught:
        extract_document("scan.pdf", b"this is not a pdf")
    assert caught.value.status_code == 400


def test_upload_validation():
    payload = validate_upload("Quarterly Report.pdf", b"%PDF", max_bytes=100)
    assert payload.filename == "Quarterly Report.pdf"
    assert payload.media_type == "application/pdf"

    with pytest.raises(AppError) as bad_type:
        validate_upload("notes.txt", b"hello", max_bytes=100)
    assert bad_type.value.status_code == 400

    with pytest.raises(AppError) as empty:
        validate_upload("notes.docx", b"", max_bytes=100)
    assert empty.value.status_code == 400

    with pytest.raises(AppError) as large:
        validate_upload("notes.pdf", b"123456", max_bytes=4)
    assert large.value.status_code == 413


def test_titan_request_uses_configured_dimensions():
    assert titan_embed_request("hello", 1024) == {
        "inputText": "hello",
        "dimensions": 1024,
        "normalize": True,
    }


def test_bedrock_response_text_joins_content_blocks():
    payload = {"output": {"message": {"content": [{"text": "Answer [1]."}, {"text": "More."}]}}}
    assert _response_text(payload) == "Answer [1].\nMore."
    assert _response_text({"output": {}}) == ""
    json.dumps(payload)
